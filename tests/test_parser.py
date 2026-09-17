"""test_parser.py — контракт поиска, форматы Excel и атомарность обновления."""

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import xlwt

from btdetal import BtdetalError, BtdetalParser, PriceImportError

HEADERS = ["Код", "Артикул", "Номенклатура", "Остаток", "1-я цена"]
ROWS = [
    ["00001", "00123", "Втулка Braun MR65", 8, 62],
    ["00002", "BR67050811", "Втулка без остатка", 0, 60],
    ["00003", None, "Крышка Moulinex", 2, "1 234,56"],
    ["00004", "00123", "Втулка аналог", 3, 63],
    [None, None, "Раздел", 0, None],
]


def workbook(rows: list[list[Any]] | None = None, *, kind: str = "xlsx") -> bytes:
    """Создать минимальный тестовый файл без данных реального поставщика."""
    output = BytesIO()
    values = [HEADERS, *(ROWS if rows is None else rows)]
    if kind == "xls":
        book = xlwt.Workbook()
        sheet = book.add_sheet("Прайс")
        for r, row in enumerate(values):
            for c, value in enumerate(row):
                sheet.write(r, c, value)
        book.save(output)
    else:
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        for row in values:
            sheet.append(row)
        book.save(output)
        book.close()
    return output.getvalue()


@pytest.fixture
def parser(tmp_path: Path) -> BtdetalParser:
    """Подготовить независимую заполненную базу."""
    client = BtdetalParser(tmp_path / "price.sqlite3")
    _ = client.import_price(workbook(), price_date="2026-09-14")
    return client


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_formats(tmp_path: Path, kind: str) -> None:
    """Оба Excel-формата дают одинаковые цены, остатки и ведущие нули."""
    parser = BtdetalParser(tmp_path / "price.sqlite3")
    report = parser.import_price(BytesIO(workbook(kind=kind)))
    assert report["imported_count"] == 3
    assert report["skipped_zero_stock"] == 1
    assert report["skipped_sections"] == 1
    result = parser.search("00123")
    assert result["count"] == 2
    product = result["products"][0]
    assert product["id"] == "00001"
    assert product["price"] == "62.00"
    assert product["currency"] == "RUB"
    assert product["stock_quantity"] == 8
    assert product["exact_match"] is True
    assert product["quantity_prices"] == []
    assert product["url"] is None
    assert parser.search("BR67050811")["status"] == "not_found"
    found = parser.search("КРЫШКА")["products"][0]
    assert found["article"] == ""
    assert found["price"] == "1234.56"
    assert found["exact_match"] is False


def test_limit_and_batch(parser: BtdetalParser) -> None:
    """Широкий поиск не усечён; пакет сохраняет порядок и повторные запросы."""
    result = parser.search("втулка", max_results=1)
    assert result["status"] == "too_many_results"
    assert result["products"] == []
    assert result["count_at_least"] == 2
    batch = parser.search_many(["00123", "нет товара", "00123"], max_workers=3)
    assert [r["query"] for r in batch["results"]] == ["00123", "нет товара", "00123"]
    assert batch["summary"]["ok"] == 2
    assert batch["summary"]["not_found"] == 1


@pytest.mark.parametrize("query", ["вт", "Braun", "MR65", "00123", "КРЫШКА"])
def test_substrings(parser: BtdetalParser, query: str) -> None:
    """Находить в обеих колонках, включая кириллицу и короткие запросы."""
    assert parser.search(query)["status"] == "ok"


@pytest.mark.parametrize("query", ['a"b', "a%b", "a_b", "OR", "--", "*", "%", "_"])
def test_literal_queries(parser: BtdetalParser, query: str) -> None:
    """Спецсимволы SQL/FTS трактуются как буквальный текст."""
    _ = parser.import_price(workbook([["a", query, "Запчасть", 1, 1]]))
    assert parser.search(query)["count"] == 1
    assert parser.search("' OR 1=1 --")["status"] == "not_found"


@pytest.mark.parametrize(
    "stock,price", [(None, 3), (-1, 3), ("много", 3), (1, "NaN"), (1, -4), (True, 2)]
)
def test_failed_import_preserves_database(parser: BtdetalParser, stock: Any, price: Any) -> None:
    """Неверные числа не обнуляют и не частично обновляют действующий прайс."""
    previous = parser.get_metadata()
    with pytest.raises(PriceImportError):
        _ = parser.import_price(workbook([["x", "a", "Товар", stock, price]]))
    assert parser.get_metadata() == previous
    assert parser.search("00123")["count"] == 2


def test_replace_and_empty(parser: BtdetalParser) -> None:
    """Обновление удаляет устаревшие позиции; пустая база требует явного выбора."""
    with pytest.raises(PriceImportError):
        _ = parser.import_price(workbook([]))
    _ = parser.import_price(workbook([["b", "новый", "Товар", 1, 10]]))
    assert parser.search("00123")["status"] == "not_found"
    assert parser.search("новый")["count"] == 1
    _ = parser.import_price(workbook([]), allow_empty=True)
    assert parser.search("новый")["status"] == "not_found"


def test_duplicate_code(parser: BtdetalParser) -> None:
    """Совпадающие артикулы допустимы, совпадающие коды отклоняются."""
    with pytest.raises(PriceImportError, match="Повтор кода"):
        _ = parser.import_price(workbook([ROWS[0], ROWS[0]]))


@pytest.mark.parametrize("kind", ["xlsx", "xls"])
def test_numeric_identifiers(tmp_path: Path, kind: str) -> None:
    """Сохранять ведущие нули, заданные числовым форматом Excel."""
    stream = BytesIO()
    if kind == "xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append(HEADERS)
        sheet.append([21, 123, "Деталь", 1, 4])
        sheet["A2"].number_format = "00000000000;[Red]\\-00000000000"
        sheet["B2"].number_format = "00000"
        book.save(stream)
    else:
        book = xlwt.Workbook()
        sheet = book.add_sheet("Прайс")
        for c, label in enumerate(HEADERS):
            sheet.write(0, c, label)
        for c, value in enumerate([21, 123, "Деталь", 1, 4]):
            style = xlwt.easyxf(num_format_str="00000000000" if c == 0 else "00000")
            sheet.write(1, c, value, style)
        book.save(stream)
    parser = BtdetalParser(tmp_path / "db.sqlite3")
    _ = parser.import_price(stream.getvalue())
    assert parser.search("00123")["products"][0]["id"] == "00000000021"


def test_path_and_restart(tmp_path: Path) -> None:
    """Файл принимается по пути, а база доступна новому экземпляру без Excel."""
    file = tmp_path / "price.xlsx"
    _ = file.write_bytes(workbook())
    db = tmp_path / "db.sqlite3"
    _ = BtdetalParser(db).import_price(file)
    file.unlink()
    parser = BtdetalParser(db)
    assert parser.get_product("00001")["article"] == "00123"
    assert parser.get_metadata()["source_name"] == "price.xlsx"
    with pytest.raises(KeyError):
        _ = parser.get_product("missing")


def test_uninitialized(tmp_path: Path) -> None:
    """Отсутствующий прайс не маскируется под отсутствие товара."""
    parser = BtdetalParser(tmp_path / "absent.sqlite3")
    with pytest.raises(BtdetalError):
        _ = parser.search("00123")
    assert not parser.db_path.exists()


def test_concurrent_readers(parser: BtdetalParser) -> None:
    """Один экземпляр допускает обращения из разных потоков."""
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(parser.search, ["00123"] * 20))
    assert all(result["count"] == 2 for result in results)


def test_reader_snapshot_during_import(parser: BtdetalParser) -> None:
    """Открытая транзакция видит прежнюю целую версию, новая — обновлённую."""
    import sqlite3

    conn = sqlite3.connect(parser.db_path)
    try:
        _ = conn.execute("BEGIN")
        assert conn.execute("SELECT count(*) FROM products").fetchone()[0] == 3
        _ = parser.import_price(workbook([["b", "новый", "Товар", 1, 10]]))
        assert conn.execute("SELECT count(*) FROM products").fetchone()[0] == 3
    finally:
        conn.close()
    assert parser.search("новый")["count"] == 1


def test_sheet_selection_and_formulas(parser: BtdetalParser) -> None:
    """Не выбирать произвольный лист и не превращать формулу без кеша в ноль."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(HEADERS)
    sheet.append(ROWS[0])
    _ = book.copy_worksheet(sheet)
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(PriceImportError):
        _ = parser.import_price(stream.getvalue())
    assert parser.import_price(stream.getvalue(), sheet_name=sheet.title)["imported_count"] == 1
    sheet["E2"] = "=1+2"
    stream = BytesIO()
    book.save(stream)
    report = parser.import_price(stream.getvalue(), sheet_name=sheet.title)
    assert report["missing_price_count"] == 1
    assert parser.search("00123")["status"] == "partial_error"


@pytest.mark.parametrize("query", ["", " ", "a\x00b", None, 123])
def test_invalid_query(parser: BtdetalParser, query: Any) -> None:
    """Отклонять некорректный запрос и пакет до обращения к базе."""
    with pytest.raises(ValueError):
        _ = parser.search(query)
    with pytest.raises(ValueError):
        _ = parser.search_many(["00123", query])


@pytest.mark.parametrize("limit", [0, -1, True, 1.5])
def test_invalid_limit(parser: BtdetalParser, limit: Any) -> None:
    """Лимиты должны быть положительными целыми."""
    with pytest.raises(ValueError):
        _ = parser.search("00123", max_results=limit)


def test_decimal_article(parser: BtdetalParser) -> None:
    """Артикул с точкой может храниться числом и не должен ломать импорт."""
    _ = parser.import_price(workbook([["a", 651017674.5, "Модуль", 1, 2718]]))
    assert parser.search("651017674.5")["products"][0]["article"] == "651017674.5"


def test_database_write_rollback(parser: BtdetalParser) -> None:
    """Ошибка записи после удаления старых строк откатывает весь импорт."""
    import sqlite3

    previous = parser.get_metadata()
    conn = sqlite3.connect(parser.db_path)
    try:
        _ = conn.execute("""CREATE TRIGGER block_import BEFORE INSERT ON products
            BEGIN SELECT RAISE(ABORT, 'test failure'); END""")
        conn.commit()
        with pytest.raises(BtdetalError):
            _ = parser.import_price(workbook([["new", "new", "Товар", 1, 3]]))
    finally:
        conn.close()
    assert parser.search("00123")["count"] == 2
    assert parser.get_metadata() == previous


def test_corrupted_file(parser: BtdetalParser) -> None:
    """Повреждённый файл и отсутствующие заголовки не заменяют прайс."""
    previous = parser.get_metadata()
    for data in (b"not an excel file", b"PKbroken"):
        with pytest.raises(PriceImportError):
            _ = parser.import_price(data)
    book = openpyxl.Workbook()
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(PriceImportError):
        _ = parser.import_price(stream.getvalue())
    assert parser.get_metadata() == previous


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_missing_price(parser: BtdetalParser, kind: str) -> None:
    """Товар без цены остаётся в наличии с явной ошибкой, остальные доступны."""
    report = parser.import_price(
        workbook([*ROWS, ["missing", "belt", "Ремень", 2, None]], kind=kind)
    )
    assert report["missing_price_count"] == 1
    result = parser.search("belt")
    assert result["status"] == "partial_error"
    product = result["products"][0]
    assert product["price"] is None
    assert product["raw_price"] is None
    assert product["status"] == "error"
    assert product["price_is_current"] is False
    assert product["stock_quantity"] == 2
    batch = parser.search_many(["belt", "00123"])
    assert batch["status"] == "partial_error"
    assert batch["summary"]["partial_error"] == 1
    assert batch["summary"]["ok"] == 1
